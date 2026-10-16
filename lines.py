"""
Goldberg's lines.

Every quip lives in the "Lines" tab of a Google Sheet (columns: Category, Line, Active):
the one named by LINES_SHEET_KEY, or the team's sheet (DATA_SHEET_KEY) if that's unset.
If the sheet can't be reached, the bundled goldberg_lines.xlsx is used instead.

Lines are picked from a shuffled pile per category, so nothing repeats until every
active line in that category has been used.
"""

import os
import random
import time

LINES_WORKSHEET = "Lines"
LOCAL_LINES_FILE = os.path.join(os.path.dirname(__file__), "goldberg_lines.xlsx")
REFRESH_SECONDS = 300

_lines: dict[str, list[str]] = {}
_loaded_at: float | None = None
_piles: dict[str, list[str]] = {}
_last_said: dict[str, str] = {}


class _KeepUnknown(dict):
    """Leaves unknown {placeholders} as-is instead of raising."""
    def __missing__(self, key):
        return "{" + key + "}"


def _is_active(value) -> bool:
    if value is None or value == "":
        return True
    return str(value).strip().upper() not in ("FALSE", "NO", "0")


def _parse_rows(rows: list) -> dict[str, list[str]]:
    if not rows:
        return {}
    headers = [str(h or "").strip().lower() for h in rows[0]]
    cat_i, line_i = headers.index("category"), headers.index("line")
    active_i = headers.index("active") if "active" in headers else None

    lines: dict[str, list[str]] = {}
    for row in rows[1:]:
        row = list(row) + [None] * (len(headers) - len(row))
        category = str(row[cat_i] or "").strip()
        line = str(row[line_i] or "").strip()
        if not category or not line:
            continue
        if active_i is not None and not _is_active(row[active_i]):
            continue
        lines.setdefault(category, []).append(line)
    return lines


def _load_from_sheet() -> dict[str, list[str]]:
    from config import LINES_SHEET_KEY
    from utils import get_worksheet
    return _parse_rows(get_worksheet(LINES_WORKSHEET, LINES_SHEET_KEY).get_all_values())


def _load_from_file() -> dict[str, list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(LOCAL_LINES_FILE, read_only=True, data_only=True)
    try:
        return _parse_rows(list(wb[LINES_WORKSHEET].iter_rows(values_only=True)))
    finally:
        wb.close()


def reload_lines() -> tuple[int, str]:
    """Re-read the lines now. Returns (line count, source)."""
    global _lines, _loaded_at
    _loaded_at = time.monotonic()

    loaded, source = None, None
    for loader, name in ((_load_from_sheet, "Google Sheet"), (_load_from_file, "goldberg_lines.xlsx")):
        try:
            loaded = loader()
        except Exception as e:
            print(f"[Lines] ⚠️  Couldn't load lines from {name}: {e}")
            continue
        if loaded:
            source = name
            break

    if not loaded:
        return sum(len(v) for v in _lines.values()), "previous copy (nothing could be loaded)"

    if loaded != _lines:
        _lines = loaded
        _piles.clear()
    count = sum(len(v) for v in _lines.values())
    print(f"[Lines] ✅ Loaded {count} lines in {len(_lines)} categories from {source}")
    return count, source


def get_line(category: str, **values) -> str:
    """Pick a line from `category` and fill in any {placeholders} from `values`."""
    if _loaded_at is None or time.monotonic() - _loaded_at > REFRESH_SECONDS:
        reload_lines()

    pool = _lines.get(category)
    if not pool:
        print(f"[Lines] ⚠️  No active lines for category '{category}'")
        return "..."

    pile = _piles.get(category)
    if not pile:
        pile = pool.copy()
        random.shuffle(pile)
        # Don't let a fresh pile start with the line we just said
        if len(pile) > 1 and pile[-1] == _last_said.get(category):
            pile[0], pile[-1] = pile[-1], pile[0]
        _piles[category] = pile

    line = pile.pop()
    _last_said[category] = line
    try:
        return line.format_map(_KeepUnknown(values))
    except (ValueError, IndexError):
        return line
